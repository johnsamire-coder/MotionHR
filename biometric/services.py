"""
منطق البصمة المشترك: قراية البصمات (من الجهاز أو من ملف)، ربطها بالموظفين، وحفظها من غير تكرار.
"""
import csv
import io
import re
from datetime import datetime

from django.db import transaction
from django.utils import timezone

from biometric.models import BiometricPunch, BiometricUserMap

DATE_FORMATS = ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y/%m/%d %H:%M:%S', '%Y/%m/%d %H:%M',
                '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M',
                '%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y %I:%M %p', '%d/%m/%Y %I:%M:%S %p', '%d/%m/%Y %I:%M %p')


# ─────────────────────────── الوقت ───────────────────────────

def parse_time(value):
    """وقت الجهاز بيبقى بتوقيت مصر من غير منطقة زمنية؛ بنحوله لوقت كامل"""
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value or '').strip()
        dt = None
        for fmt in DATE_FORMATS:
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt.replace(microsecond=0)


def clean_pin(value):
    text = str(value or '').strip()
    if re.fullmatch(r'\d+\.0', text):      # أرقام الإكسل بتيجي 15.0
        text = text[:-2]
    return text[:32]


# ─────────────────────────── ربط رقم الجهاز بالموظف ───────────────────────────

def resolve_employees(company_id, pins):
    """
    {pin: employee_id أو None}
    1) الربط اليدوي اللي اتعمل من شاشة الربط
    2) رقم الجهاز = كود الموظف بالظبط
    3) رقم الجهاز = الأرقام اللي في كود الموظف (EMP00015 = 15) لو موظف واحد بس
    """
    from employees.models import Employee
    pins = {p for p in pins if p}
    out = {p: None for p in pins}
    if not company_id or not pins:
        return out
    for m in BiometricUserMap.objects.filter(company_id=company_id, pin__in=pins).exclude(employee_id=None):
        out[m.pin] = m.employee_id
    left = {p for p in pins if out[p] is None}
    if not left:
        return out
    by_code, by_digits = {}, {}
    for eid, code in Employee._base_manager.filter(company_id=company_id).values_list('id', 'employee_code'):
        code = (code or '').strip()
        if not code:
            continue
        by_code.setdefault(code.lower(), []).append(eid)
        digits = re.sub(r'\D', '', code).lstrip('0')
        if digits:
            by_digits.setdefault(digits, []).append(eid)
    for p in left:
        exact = by_code.get(p.lower(), [])
        if len(exact) == 1:
            out[p] = exact[0]
            continue
        nums = by_digits.get(p.lstrip('0') or '0', [])
        if len(nums) == 1:
            out[p] = nums[0]
    return out


# ─────────────────────────── الحفظ ───────────────────────────

def store_punches(company_id, device_serial, rows, source='device'):
    """
    rows: [{'pin', 'time', 'status', 'verify'}]
    بترجع (عدد الجديد، عدد المكرر، عدد اللي ماتربطش بموظف)
    """
    rows = [r for r in rows if r.get('pin') and r.get('time')]
    if not rows:
        return 0, 0, 0
    mapping = resolve_employees(company_id, {r['pin'] for r in rows})
    objs = [BiometricPunch(
        company_id=company_id, device_serial=device_serial[:64], pin=r['pin'], punch_time=r['time'],
        status_code=str(r.get('status') or '')[:8], verify_code=str(r.get('verify') or '')[:8],
        employee_id=mapping.get(r['pin']), source=source,
    ) for r in rows]
    # البصمة المكررة (نفس الجهاز ونفس الموظف ونفس الثانية) مابتتسجلش مرتين
    keys = {(o.device_serial, o.pin, o.punch_time) for o in objs}
    existing = set(BiometricPunch.objects.filter(
        device_serial=device_serial[:64], pin__in={o.pin for o in objs},
        punch_time__in={o.punch_time for o in objs},
    ).values_list('device_serial', 'pin', 'punch_time'))
    fresh, seen = [], set()
    for o in objs:
        k = (o.device_serial, o.pin, o.punch_time)
        if k in existing or k in seen:
            continue
        seen.add(k)
        fresh.append(o)
    with transaction.atomic():
        BiometricPunch.objects.bulk_create(fresh, ignore_conflicts=True)
    unmatched = sum(1 for o in fresh if not o.employee_id)
    return len(fresh), len(keys) - len(fresh), unmatched


def relink_company(company_id):
    """بعد ربط رقم جهاز بموظف: البصمات القديمة اللي ماكانتش مربوطة بتتربط"""
    pending = BiometricPunch.objects.filter(company_id=company_id, employee_id=None)
    pins = set(pending.values_list('pin', flat=True).distinct())
    mapping = resolve_employees(company_id, pins)
    n = 0
    for pin, eid in mapping.items():
        if eid:
            n += pending.filter(pin=pin).update(employee_id=eid)
    return n


# ─────────────────────────── قراية بيانات الجهاز ───────────────────────────

def parse_attlog(text):
    """
    سطور ATTLOG من الجهاز (ADMS) أو من ملف attlog.dat على الفلاشة:
      PIN \\t 2026-10-21 08:59:12 \\t status \\t verify \\t workcode ...
    """
    rows = []
    for line in (text or '').splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in re.split(r'\t|\s{2,}', line) if p.strip()]
        if len(parts) < 2:
            continue
        pin = clean_pin(parts[0])
        t = parse_time(parts[1])
        if t is None and len(parts) >= 3:            # التاريخ والوقت في عمودين
            t = parse_time(f'{parts[1]} {parts[2]}')
            parts = parts[:1] + [f'{parts[1]} {parts[2]}'] + parts[3:]
        if not pin or t is None:
            continue
        rows.append({'pin': pin, 'time': t,
                     'status': parts[2] if len(parts) > 2 else '',
                     'verify': parts[3] if len(parts) > 3 else ''})
    return rows


def parse_userinfo(text):
    """سطور USERINFO من الجهاز: USER PIN=15\\tName=Ahmed\\t... → {pin: name}"""
    out = {}
    for line in (text or '').splitlines():
        fields = dict(kv.split('=', 1) for kv in line.replace('USER ', '', 1).split('\t') if '=' in kv)
        pin = clean_pin(fields.get('PIN'))
        if pin:
            out[pin] = (fields.get('Name') or '').strip()[:150]
    return out


PIN_HEADERS = ('ac-no', 'ac no', 'acno', 'no.', 'no', 'pin', 'user id', 'userid', 'id', 'enroll',
               'emp no', 'badge', 'رقم', 'الرقم', 'كود', 'رقم الموظف')
TIME_HEADERS = ('date/time', 'datetime', 'time', 'check time', 'checktime', 'punch time', 'التاريخ والوقت', 'الوقت')
DATE_HEADERS = ('date', 'التاريخ')


def _pick(headers, names):
    for i, h in enumerate(headers):
        if h in names:
            return i
    for i, h in enumerate(headers):
        if any(n in h for n in names if len(n) > 3):
            return i
    return None


def _rows_from_table(table):
    """جدول (من CSV أو إكسل): بندوّر على عمود رقم الموظف وعمود الوقت (أو التاريخ + الوقت)"""
    table = [r for r in table if any(str(c or '').strip() for c in r)]
    if not table:
        return []
    headers = [str(h or '').strip().lower() for h in table[0]]
    pin_i = _pick(headers, PIN_HEADERS)
    time_i = _pick(headers, TIME_HEADERS)
    date_i = _pick(headers, DATE_HEADERS)
    if pin_i is None or (time_i is None and date_i is None):
        return []
    rows = []
    for r in table[1:]:
        r = list(r) + [''] * len(headers)
        pin = clean_pin(r[pin_i])
        tv = r[time_i] if time_i is not None else None
        if date_i is not None and time_i is not None and date_i != time_i and not isinstance(tv, datetime):
            dv = r[date_i]
            dv = dv.strftime('%Y-%m-%d') if isinstance(dv, datetime) else str(dv or '').strip()
            tv = tv.strftime('%H:%M:%S') if hasattr(tv, 'strftime') else str(tv or '').strip()
            tv = f'{dv} {tv}'
        elif time_i is None:
            tv = r[date_i]
        t = parse_time(tv)
        if pin and t:
            rows.append({'pin': pin, 'time': t, 'status': '', 'verify': ''})
    return rows


def parse_upload(name, data):
    """ملف من الفلاشة أو من برنامج الجهاز: attlog.dat / txt / csv / xlsx"""
    name = (name or '').lower()
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            rows += _rows_from_table(list(ws.iter_rows(values_only=True)))
        return rows
    text = None
    for enc in ('utf-8-sig', 'utf-16', 'cp1256', 'latin-1'):
        try:
            text = data.decode(enc)
            if '\x00' not in text:
                break
        except UnicodeDecodeError:
            continue
    text = text or ''
    if name.endswith('.csv'):
        return _rows_from_table(list(csv.reader(io.StringIO(text))))
    rows = parse_attlog(text)
    if not rows:     # ملف txt متصدّر كجدول بعناوين
        rows = _rows_from_table([re.split(r'\t|,', l) for l in text.splitlines()])
    return rows
